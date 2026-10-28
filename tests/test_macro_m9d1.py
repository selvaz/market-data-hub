"""Offline official Japanese/Chinese publication and edition regression tests."""
from dataclasses import replace
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import pytest
import requests

from market_data_hub.macro_releases import open_writer, run_daily
from market_data_hub.macro_releases.adapters import persist
from market_data_hub.macro_releases.asia import AsiaAdapter
from market_data_hub.macro_releases.asia_editions import EDITIONS, SCHEDULES
from market_data_hub.macro_releases.catalogue import ASIA_FAMILIES
from market_data_hub.macro_releases.http import Document, FetchError
from market_data_hub.macro_releases.schedule import RUN_INSTANT, due

ROOT = Path(__file__).parent / "fixtures/macro_releases/m9d1"
GOLDEN = json.loads((ROOT / "golden-set.json").read_text())["cases"]


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Network forbidden in M9d1 tests")
    monkeypatch.setattr(requests.Session, "request", forbidden)
    token = RUN_INSTANT.set("2026-10-27T18:00:00Z")
    yield
    RUN_INSTANT.reset(token)


def document(name):
    body = (ROOT / name).read_bytes()
    meta = json.loads((ROOT / (name + ".meta.json")).read_text())
    assert hashlib.sha256(body).hexdigest() == meta["document_hash"]
    return Document(body, meta["url"], meta["retrieved_at"], meta["document_hash"])


def parser(family):
    from market_data_hub.macro_releases import nbs_macro, japan_macro
    return (japan_macro if family.startswith("jp_") else nbs_macro).parse


@pytest.mark.parametrize("case", GOLDEN, ids=lambda c: c["family"] + ":" + c["edition"] + ":" + c["indicator"])
def test_issuer_golden_and_read_card(case, tmp_path):
    rows = parser(case["family"])(document(case["fixture"]), case["family"])
    row = next(r for r in rows if r.indicator_id == case["indicator"])
    assert (row.value, row.period, row.edition, row.stage) == (
        case["value"], case["period"], case["edition"], case["stage"])
    assert row.document.url == case["source_url"]
    assert row.series_key.startswith(row.indicator_id + ":")
    with open_writer(tmp_path / "golden.sqlite") as store:
        persist(store, [row])
        utc_day = row.published_at[:10]
        cards = store.release_cards(utc_day, utc_day, indicators=[row.indicator_id])["cards"]
        actual = next(c for c in cards if c.get("resolution"))
        assert actual["source_class"] == actual["resolution"]["status"] == "official"
        assert actual["resolution"]["value"] == case["value"]
        assert actual["resolution"]["edition"] == case["edition"]


def test_family_calendars_and_golden_completeness():
    for family in ASIA_FAMILIES:
        cases = [c for c in GOLDEN if c["family"] == family]
        assert len({c["edition"] for c in cases}) >= 2
        assert SCHEDULES[family].startswith("https://")
        assert max(EDITIONS[family]) <= "2026-12-31"
        assert {c["indicator"] for c in cases} == {i.id for i in ASIA_FAMILIES[family]}


class Replay:
    def __init__(self, **kwargs):
        self.requests = []

    def get(self, provider, url, params=None):
        if params:
            from urllib.parse import urlencode
            url += "?" + urlencode(params)
        self.requests.append(url)
        for path in ROOT.glob("*.meta.json"):
            meta = json.loads(path.read_text())
            if meta["url"] == url:
                return document(path.name.removesuffix(".meta.json"))
        raise AssertionError("Unrecorded issuer URL: " + url)


@pytest.mark.parametrize("family", ASIA_FAMILIES)
def test_daily_schedules_and_idempotence(family, tmp_path, monkeypatch):
    monkeypatch.setattr(run_daily, "HttpClient", Replay)
    class Clock:
        @staticmethod
        def now(tz):
            return datetime(2026, 10, 27, 18, tzinfo=timezone.utc)
    monkeypatch.setattr(run_daily, "datetime", Clock)
    target = tmp_path / "daily.sqlite"
    first = run_daily.execute(target, [family], 24, False, raw_dir=tmp_path / "none")
    assert first["families"][family]["status"] == "ok", first
    assert first["families"][family]["vintages_added"] >= 2
    second = run_daily.execute(target, [family], 24, False, raw_dir=tmp_path / "none")
    assert second["families"][family]["status"] == "no_change", second


def test_cpi_unknown_date_clock_measure_and_disagreement_fail_closed():
    doc = document("cn-cpi-sep.html")
    for before, after in ((b"2026/10/15 09:30", b"2026/10/16 09:30"),
                          (b"2026/10/15 09:30", b"2026/10/15 10:00"),
                          (b"Consumer Price Index in September 2026", b"Consumer Price Index in August 2026"),
                          (b"increased by 0.9% year on year", b"increased by 9.9% year on year")):
        assert before in doc.body
        with pytest.raises(FetchError, match="needs review"):
            parser("cn_cpi")(replace(doc, body=doc.body.replace(before, after)), "cn_cpi")


def test_release_instant_and_horizon_fail_closed():
    client = Replay()
    token = RUN_INSTANT.set("2026-09-10T01:29:59Z")
    try:
        assert AsiaAdapter(client).fetch(["cn_cpi"], "2026-01") == []
        assert client.requests == []
    finally:
        RUN_INSTANT.reset(token)
    token = RUN_INSTANT.set("2027-01-01T00:00:00Z")
    try:
        with pytest.raises(FetchError, match="needs review"):
            AsiaAdapter(client).fetch(["cn_cpi"], "2026-01")
    finally:
        RUN_INSTANT.reset(token)


@pytest.mark.parametrize("family", ASIA_FAMILIES)
def test_unreviewed_source_url_fails_closed(family):
    case = next(c for c in GOLDEN if c["family"] == family)
    with pytest.raises(FetchError, match="needs review"):
        parser(family)(replace(document(case["fixture"]), url="https://example.invalid/unreviewed"), family)


def test_gdp_cumulative_columns_and_constant_price_basis_are_required():
    doc = document("cn-gdp-native-q3.html")
    for before, after in (("前三季度", "三季度"), ("增长速度按不变价计算", "增长速度按现价计算")):
        raw = doc.body.decode("utf-8")
        assert before in raw
        with pytest.raises(FetchError, match="needs review"):
            parser("cn_gdp")(replace(doc, body=raw.replace(before, after).encode()), "cn_gdp")


def test_tokyo_release_date_claim_and_local_release_instant():
    assert "2026-10-22" not in EDITIONS["jp_tokyo_cpi"]
    assert EDITIONS["jp_tokyo_cpi"]["2026-10-30"][0] == "2026-10"
    client = Replay()
    token = RUN_INSTANT.set("2026-10-29T23:29:59Z")
    try:
        assert AsiaAdapter(client).fetch(["jp_tokyo_cpi"], "2026-10") == []
        assert client.requests == []
    finally:
        RUN_INSTANT.reset(token)
    assert "2026-10-30" in due(EDITIONS["jp_tokyo_cpi"], "2026-10-29T23:30:00Z", family="jp_tokyo_cpi")


def test_only_reviewed_stages_have_calendar_membership():
    from market_data_hub.macro_releases.read_api import indicators_for_stage
    for family in ASIA_FAMILIES:
        assert indicators_for_stage(family, "unreviewed-stage") == set()
    assert indicators_for_stage("jp_trade_balance", "final") == set()
    assert indicators_for_stage("jp_tokyo_cpi", "final") == set()


def test_real_issuer_schedule_evidence_covers_the_reviewed_horizon():
    from io import BytesIO
    from openpyxl import load_workbook
    from market_data_hub.macro_releases.tier1_publications import publication_text
    files = {json.loads(p.read_text())["url"] for p in ROOT.glob("*.meta.json")}
    assert set(SCHEDULES.values()).issubset(files)
    cpi, _ = publication_text(document("cpi-calendar.html"))
    assert "October 30" in cpi and "December 25" in cpi and "December 18" in cpi
    nbs, _ = publication_text(document("nbs-calendar.html"))
    assert "Regular Press Release Calendar of NBS in 2026" in nbs
    assert "31/Sat 30/Mon 31/Thu" in nbs
    schedule_pdf = document("jp-gdp-schedule.pdf")
    schedule_meta = json.loads((ROOT / "jp-gdp-schedule.txt.meta.json").read_text())
    assert schedule_meta["source_document_hash"] == schedule_pdf.document_hash
    gdp = document("jp-gdp-schedule.txt").body.decode("utf-8")
    assert "November 16, 2026" in gdp and "December 8, 2026" in gdp
    workbook = load_workbook(BytesIO(document("boj-release.xlsx").body), data_only=True)
    try:
        rows = [r for r in workbook['Statistics data'].values
                if r[1] == 'Tankan (Short-Term Economic Survey of Enterprises in Japan) / Summary and Outline']
        assert rows[0][3] == '8:50 a.m.'
        assert rows[1][7].date().isoformat() == '2026-10-01'
        assert rows[1][9].date().isoformat() == '2026-12-14'
    finally:
        workbook.close()


def test_e_stat_index_sheet_cannot_replace_published_rate_sheet():
    from io import BytesIO
    from openpyxl import load_workbook
    doc = document("jp-cpi-sep.xlsx")
    workbook = load_workbook(BytesIO(doc.body))
    try:
        sheet = next(s for s in workbook if s['H8'].value == 'Monthly, % Change over the year')
        sheet['H8'] = 'Monthly, Index'
        stream = BytesIO()
        workbook.save(stream)
    finally:
        workbook.close()
    with pytest.raises(FetchError, match="needs review"):
        parser("jp_cpi")(replace(doc, body=stream.getvalue()), "jp_cpi")


def test_trade_units_and_provisional_stage_must_match():
    from market_data_hub.macro_releases import japan_macro
    style = document("trade-table.xsl")
    for before, after in ((b"millions of YEN", b"billions of YEN"),
                          (b"not seasonally adjusted", b"seasonally adjusted")):
        assert before in style.body
        with pytest.raises(FetchError, match="needs review"):
            japan_macro.trade_stylesheet(replace(style, body=style.body.replace(before, after)))
    doc = document("trade-sep.xml")
    for before, after in ((b"(Provisional)", b"(Final)"),
                          (b"<sashihikikagakue>-536,558", b"<sashihikikagakue>-536,559"),
                          (b"<chiikikuni>Grand Total", b"<chiikikuni>ASIA")):
        assert before in doc.body
        with pytest.raises(FetchError, match="needs review"):
            japan_macro.parse(replace(doc, body=doc.body.replace(before, after)), "jp_trade_balance")


def test_esri_csv_requires_printed_rate_headers_and_latest_quarter():
    doc = document("jp-gdp-first.csv")
    for before, after in ((b"Quarter-to-Quarter Percent Change", b"Annualized Percent Change"),
                          (b"GDP(Expenditure Approach)", b"PrivateConsumption"),
                          (b"2026/ 1- 3.", b"2027/ 1- 3.")):
        assert before in doc.body
        with pytest.raises(FetchError, match="needs review"):
            parser("jp_gdp")(replace(doc, body=doc.body.replace(before, after)), "jp_gdp")
    for body in (b",\n" * 10, b"invalid encoding: \x81"):
        with pytest.raises(FetchError, match="needs review"):
            parser("jp_gdp")(replace(doc, body=body), "jp_gdp")


def test_esri_csv_release_date_is_bound_to_the_official_archive():
    from market_data_hub.macro_releases.japan_macro import gdp_archive_listing
    archive = document("jp-gdp-archive.html")
    menu = document("jp-gdp-first-menu.html")
    gdp_archive_listing(archive, menu, "2026-08-17")
    for before, after in ((b"Aug  17, 2026", b"Aug  18, 2026"),
                          (b"First preliminary Estimates", b"Second preliminary Estimates")):
        assert before in archive.body
        with pytest.raises(FetchError, match="needs review"):
            gdp_archive_listing(replace(archive, body=archive.body.replace(before, after)), menu, "2026-08-17")


def test_japan_adapters_do_not_require_optional_pdf_aes(monkeypatch):
    # Minimal installations omit both pypdf AES backends. Only Tankan's
    # unencrypted PDF should reach PdfReader in the revised adapters.
    import pypdf
    def forbidden(*args, **kwargs):
        raise AssertionError("Encrypted PDF is not a runtime dependency")
    monkeypatch.setattr(pypdf, "PdfReader", forbidden)
    for family in ("jp_gdp", "jp_trade_balance"):
        rows = AsiaAdapter(Replay()).fetch([family], "2026-01")
        assert len(rows) == 2


class ListingReplay(Replay):
    def get(self, provider, url, params=None):
        from market_data_hub.macro_releases.asia_discovery import CPI_INDEX
        # The saved index links to this exact issuer month listing.
        if url.startswith(CPI_INDEX.split('?')[0]) and 'month=23070909' in url:
            return replace(document("cpi-month-listing.html"), url=url)
        return super().get(provider, url, params)


@pytest.mark.parametrize("family,edition", [
    ("jp_trade_balance", "2026-10-29"), ("jp_tokyo_cpi", "2026-10-30"),
    ("jp_cpi", "2026-11-20"), ("jp_gdp", "2026-11-16"),
    ("cn_nbs_pmi", "2026-10-31"), ("cn_cpi", "2026-11-09"),
    ("cn_activity", "2026-11-16"),
])
def test_due_unlisted_edition_preserves_other_editions(family, edition):
    from market_data_hub.macro_releases.schedule import PENDING_EDITIONS
    instant = RUN_INSTANT.set(edition + "T18:00:00Z")
    notices = {}
    token = PENDING_EDITIONS.set(notices)
    try:
        rows = AsiaAdapter(ListingReplay()).fetch([family], "2026-01")
        assert len(rows) >= 2
        assert all(r.edition < edition for r in rows)
        assert notices[family, edition] == "scheduled_unreviewed"
    finally:
        PENDING_EDITIONS.reset(token)
        RUN_INSTANT.reset(instant)


def test_daily_reports_unreviewed_editions_without_losing_families(tmp_path, monkeypatch):
    monkeypatch.setattr(run_daily, "HttpClient", ListingReplay)
    class Clock:
        @staticmethod
        def now(tz):
            return datetime(2026, 10, 31, 18, tzinfo=timezone.utc)
    monkeypatch.setattr(run_daily, "datetime", Clock)
    result = run_daily.execute(tmp_path / "future.sqlite", ["jp_trade_balance", "cn_nbs_pmi"], 24, False,
                               raw_dir=tmp_path / "none")
    assert result["status"] == "ok"
    for family in result["families"].values():
        assert family["status"] == "ok"
        assert family["vintages_added"] == 2
        assert any(e["status"] == "scheduled_unreviewed" for e in family["scheduled_editions"])


def test_trade_discovers_a_published_url_without_reviewed_archive(monkeypatch):
    from market_data_hub.macro_releases import japan_macro
    registry = dict(japan_macro.ARCHIVES["jp_trade_balance"])
    del registry["2026-10-21"]
    monkeypatch.setitem(japan_macro.ARCHIVES, "jp_trade_balance", registry)
    rows = AsiaAdapter(ListingReplay()).fetch(["jp_trade_balance"], "2026-01")
    assert [(r.edition, r.value) for r in rows] == [("2026-09-29", -1111.913), ("2026-10-21", -536.558)]


@pytest.mark.parametrize("family,edition,count", [
    ("jp_cpi", "2026-10-23", 4), ("cn_nbs_pmi", "2026-10-01", 2),
])
def test_e_stat_and_nbs_discover_non_hardcoded_publication(monkeypatch, family, edition, count):
    from market_data_hub.macro_releases import japan_macro, nbs_macro
    from market_data_hub.macro_releases.asia_discovery import NBS_INDEX
    module = japan_macro if family.startswith("jp_") else nbs_macro
    registry = dict(module.ARCHIVES[family])
    del registry[edition]
    monkeypatch.setitem(module.ARCHIVES, family, registry)
    class Client(ListingReplay):
        def get(self, provider, url, params=None):
            if url == NBS_INDEX:
                return replace(document("nbs-prev.html"), url=url)
            return super().get(provider, url, params)
    rows = AsiaAdapter(Client()).fetch([family], "2026-01")
    assert len(rows) == count
    assert any(r.edition == edition for r in rows)


class NewTradeReplay(ListingReplay):
    mode = "ok"

    def get(self, provider, url, params=None):
        from market_data_hub.macro_releases.asia_discovery import MOF_INDEX
        if url == MOF_INDEX:
            doc = document("trade-listing.html")
            body = doc.body.replace(b"Oct.21.2026(", b"Oct.29.2026(").replace(
                b"<span>Prov.</span>", b"<span>D&amp;P</span>").replace(
                b"2026094e.xml", b"2026095e.xml")
            return replace(doc, body=body, document_hash=hashlib.sha256(body).hexdigest())
        if url.endswith("2026095e.xml"):
            if self.mode == "404":
                raise FetchError("mof: HTTP 404")
            doc = document("trade-sep.xml")
            body = doc.body.replace(b"October 21, 2026", b"October 29, 2026").replace(b"(Provisional)",
                b"(Exports:Detailed Data;Imports:Provisional(HS-code Basis))")
            if self.mode == "wrong_metadata":
                body = body.replace(b"September 2026", b"August 2026")
            return replace(doc, url=url, body=body, document_hash=hashlib.sha256(body).hexdigest())
        return super().get(provider, url, params)


@pytest.mark.parametrize("mode,instant,notice", [
    ("ok", "2026-10-29T01:00:00Z", None),
    ("wrong_metadata", "2026-10-29T01:00:00Z", "needs_review"),
    ("404", "2026-10-29T01:00:00Z", "not_yet_published"),
    ("404", "2026-10-29T06:30:00Z", "needs_review"),
])
def test_discovered_trade_validates_metadata_and_retains_404_grace(mode, instant, notice):
    from market_data_hub.macro_releases.schedule import PENDING_EDITIONS
    client = NewTradeReplay()
    client.mode = mode
    clock = RUN_INSTANT.set(instant)
    notices = {}
    token = PENDING_EDITIONS.set(notices)
    try:
        rows = AsiaAdapter(client).fetch(["jp_trade_balance"], "2026-01")
        assert len(rows) == (3 if notice is None else 2)
        assert notices.get(("jp_trade_balance", "2026-10-29")) == notice
        if notice is None:
            assert rows[-1].edition == "2026-10-29" and rows[-1].stage == "revised"
    finally:
        PENDING_EDITIONS.reset(token)
        RUN_INSTANT.reset(clock)


def test_missing_url_also_uses_publication_delay_grace():
    from market_data_hub.macro_releases.schedule import PENDING_EDITIONS
    clock = RUN_INSTANT.set("2026-10-29T01:00:00Z")
    notices = {}
    token = PENDING_EDITIONS.set(notices)
    try:
        assert len(AsiaAdapter(ListingReplay()).fetch(["jp_trade_balance"], "2026-01")) == 2
        assert notices["jp_trade_balance", "2026-10-29"] == "not_yet_published"
    finally:
        PENDING_EDITIONS.reset(token)
        RUN_INSTANT.reset(clock)


def test_quarterly_month_window_and_mof_stylesheet_binding():
    doc = document("trade-sep.xml")
    with pytest.raises(FetchError, match="needs review"):
        parser("jp_trade_balance")(replace(doc, body=doc.body.replace(b"d1201e.xsl", b"other.xsl")), "jp_trade_balance")
    for family, name in (("jp_gdp", "jp-gdp-first.csv"), ("cn_gdp", "cn-gdp-native-q3.html")):
        assert parser(family)(document(name), family, since="2026-12") == []
