"""Offline PALI units, publication gates, snapshots and public read ownership."""
from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pytest

from market_data_hub.macro_releases import open_writer, open_reader, resolve, run_daily
from market_data_hub.macro_releases import fannie_pali as pali
from market_data_hub.macro_releases.adapters import persist
from market_data_hub.macro_releases.http import FetchError, HttpClient
from market_data_hub.macro_releases.schedule import RUN_INSTANT, release_instant
from test_macro_m9f import document, offline as offline
from test_macro_housing_read_classes import assert_nonofficial


def saved():
    return document("../mortgage_applications/pali-history.xlsx")


def changed(edit):
    book = openpyxl.load_workbook(BytesIO(saved().body))
    edit(book.active)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return replace(saved(), body=buffer.getvalue())


@pytest.mark.parametrize("edit", [
    lambda sheet: setattr(sheet["B1"], "value", "WK_TOTAL_RALI_$"),
    lambda sheet: setattr(sheet["A3"], "value", "1/17/2004"),
    lambda sheet: setattr(sheet["B3"], "value", -1),
    lambda sheet: setattr(sheet["B2"], "value", 200),
    lambda sheet: sheet.delete_rows(3),
    lambda sheet: sheet.delete_rows(sheet.max_row),
    lambda sheet: setattr(sheet.cell(sheet.max_row, 2), "number_format", "0.0"),
    lambda sheet: setattr(sheet.cell(sheet.max_row - 1, 1), "value", "YoY % Change"),
])
def test_pali_format_changes_fail_closed(edit):
    with pytest.raises(FetchError, match="needs review"):
        pali.parse(changed(edit))


def test_pali_revised_history_is_not_backdated_as_original_releases():
    history, _ = pali.workbook(saved().body)
    assert history[0][0].isoformat() == "2004-01-09"
    assert len(history) == 1189
    rows = pali.parse(saved())
    assert len(rows) == 6
    assert {r.period for r in rows} == {"2026-10-16"}
    assert {r.edition for r in rows} == {"2026-10-20"}
    assert {r.published_at for r in rows} == {"2026-10-20T14:00:00+00:00"}
    assert pali.parse(saved(), since="2026-10-17") == []


class Replay:
    def __init__(self, **kwargs):
        pass

    def get(self, provider, url, params=None):
        assert (provider, url, params) == ("fanniemae", pali.URL, None)
        return saved()


def test_pali_release_instant_and_stale_rolling_snapshot():
    for instant, expected in (("2026-10-20T13:59:59+00:00", 0), ("2026-10-20T14:00:00+00:00", 6)):
        token = RUN_INSTANT.set(datetime.fromisoformat(instant))
        try:
            assert len(pali.fetch(Replay(), pali.FAMILY, "2026-10")) == expected
        finally:
            RUN_INSTANT.reset(token)
    assert release_instant("2026-09-09", family=pali.FAMILY).isoformat() == "2026-09-09T14:00:00+00:00"
    assert release_instant("2026-12-29", family=pali.FAMILY).isoformat() == "2026-12-29T15:00:00+00:00"
    token = RUN_INSTANT.set(datetime(2026, 10, 28, tzinfo=timezone.utc))
    try:
        with pytest.raises(FetchError, match="stale.*needs review"):
            pali.fetch(Replay(), pali.FAMILY, "2026-10")
    finally:
        RUN_INSTANT.reset(token)


def test_pali_collection_schedules_and_idempotence(tmp_path, monkeypatch):
    monkeypatch.setattr(run_daily, "HttpClient", Replay)
    path = tmp_path / "pali.sqlite"
    for expected in ("ok", "no_change"):
        result = run_daily.execute(path, [pali.FAMILY], 24, False, raw_dir=tmp_path / "empty")
        assert result["families"][pali.FAMILY]["status"] == expected, result
    with open_reader(path) as store:
        cards = store.release_cards("2026-10-20", "2026-10-20")["cards"]
        assert len(cards) == 6
        assert all(c["source_class"] == c["resolution"]["status"] == "official" for c in cards)
        assert all(c["issuer_id"] == "FannieMae" for c in cards)


def test_pali_mirror_cannot_become_issuer(tmp_path):
    with open_writer(tmp_path / "mirror.sqlite") as store:
        persist(store, [replace(row, source="FRED") for row in pali.parse(saved())])
        result = resolve(store, "us_fannie_pali_upb", "2026-10-16")
        assert (result.status, result.source_class) == ("provisional", "redistributor")
        assert_nonofficial(result.to_dict())


def test_public_transport_preserves_injected_http_session_and_quota():
    class Session:
        def request(self, method, url, **kwargs):
            assert method == "GET" and url == pali.URL
            return type("Response", (), dict(status_code=200, content=saved().body))()
    doc = HttpClient(session=Session(), persistent_rates=False).get("fanniemae", pali.URL)
    assert doc.body == saved().body
