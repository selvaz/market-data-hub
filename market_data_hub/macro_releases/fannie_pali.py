"""Fannie Mae's free PALI workbook; current snapshot, never historical vintages.

PALI UPB ($) and count are distinct NSA indices (2004-01-09=100).
The footer publishes fractional changes for the latest week only. History is
exposed for deterministic research, not backdated into certified releases.
"""
from datetime import datetime, timedelta
from io import BytesIO
import math

from openpyxl import load_workbook

from .housing_editions import EDITIONS, reviewed
from .http import FetchError
from .schedule import current_edition_allowed, due

FAMILY = "us_fannie_pali"
URL = "https://www.fanniemae.com/mortgage-application-index/download"
PAGE = "https://www.fanniemae.com/data-and-insights/surveys-indices/weekly-mortgage-applications-data"
CALENDAR = "https://www.fanniemae.com/media/document/pdf/fannie-mae-pali-rali-release-calendar.pdf"
HEADER = ("WK_Ending", "WK_TOTAL_PALI_$", "WK_TOTAL_PALI_CT", "WK_TOTAL_RALI_$", "WK_TOTAL_RALI_CT",
          "WK_RATE_TERM_RALI_$", "WK_RATE_TERM_RALI_CT", "WK_CASH_OUT_RALI_$", "WK_CASH_OUT_RALI_CT")


def workbook(body):
    """Validate layout, finite levels and uninterrupted Saturday-Friday weeks."""
    try:
        book = load_workbook(BytesIO(body), read_only=True, data_only=True)
        try:
            if book.sheetnames != ["Sheet1"]:
                raise ValueError()
            rows = list(book.active.values)
        finally:
            book.close()
        if tuple(rows[0]) != HEADER or len(rows) < 4:
            raise ValueError()
        history, changes = [], {}
        for row in rows[1:]:
            if len(row) != len(HEADER):
                raise ValueError()
            if row[0] in ("WoW % Change", "YoY % Change"):
                if row[0] in changes or not history:
                    raise ValueError()
                changes[row[0]] = row[1:3]
                continue
            if changes:  # Footer cannot be interspersed with observations.
                raise ValueError()
            day = datetime.strptime(row[0], "%m/%d/%Y").date()
            if day.weekday() != 4 or history and day != history[-1][0] + timedelta(days=7):
                raise ValueError()
            values = row[1:3]
            if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or v <= 0 for v in values):
                raise ValueError()
            history.append((day, *values))
        if set(changes) != {"WoW % Change", "YoY % Change"}:
            raise ValueError()
        if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or v < -1
               for values in changes.values() for v in values):
            raise ValueError()
        return history, changes
    except Exception as exc:
        if isinstance(exc, FetchError):
            raise
        raise FetchError("Fannie Mae PALI workbook format, units or weeks need review; needs review") from None


def parse(document, family=FAMILY, since="1900-01"):
    from .housing import observation
    history, changes = workbook(document.body)
    period = history[-1][0].isoformat()
    editions = [day for day, meta in EDITIONS[family].items() if meta[0] == period]
    if len(editions) != 1:
        raise FetchError("Fannie Mae PALI reference week or release calendar needs review")
    edition = editions[0]
    reviewed(family, edition)
    if period < since:
        return []
    out = []
    for column, kind in enumerate(("upb", "count")):
        values = {"": history[-1][column + 1],
                  "_ww": 100 * changes["WoW % Change"][column],
                  "_yy": 100 * changes["YoY % Change"][column]}
        for suffix, value in values.items():
            # Published levels/changes, not changes recomputed from rounded levels.
            out.append(observation(document, family, edition, f"us_fannie_pali_{kind}{suffix}",
                                   format(value, ".10g"), "xlsx"))
    return out


def fetch(client, family, since):
    if not due(EDITIONS[family], family=family):
        return []
    rows = parse(client.get("fanniemae", URL), family, since)
    if rows and not current_edition_allowed(family, rows[0].edition):
        return []
    return rows
