"""Narrow official Japan release tables, with independently printed editions."""
from datetime import datetime
from io import BytesIO, StringIO
import csv
from xml.etree import ElementTree
import re

from .asia import observation, fetch_archives
from .asia_editions import EDITIONS, reviewed
from .http import FetchError
from .tier1_publications import require, plain
from .nbs_housing import Page
from .schedule import due

ARCHIVES = {
    "jp_trade_balance": {
        "2026-09-29": "https://www.customs.go.jp/toukei/shinbun/trade-st_e/2026/2026085e.xml",
        "2026-10-21": "https://www.customs.go.jp/toukei/shinbun/trade-st_e/2026/2026094e.xml",
    },
    "jp_tankan": {
        "2026-07-01": "https://www.boj.or.jp/en/statistics/tk/gaiyo/2026/tka2606.pdf",
        "2026-10-01": "https://www.boj.or.jp/en/statistics/tk/gaiyo/2026/tka2609.pdf",
        "2026-12-14": "https://www.boj.or.jp/en/statistics/tk/gaiyo/2026/tka2612.pdf",
    },
    "jp_tokyo_cpi": {
        "2026-08-28": "https://www.e-stat.go.jp/en/stat-search/file-download?statInfId=000040491317&fileKind=4",
        "2026-10-02": "https://www.e-stat.go.jp/en/stat-search/file-download?statInfId=000040501413&fileKind=4",
    },
    "jp_cpi": {
        "2026-09-18": "https://www.e-stat.go.jp/en/stat-search/file-download?statInfId=000040501410&fileKind=4",
        "2026-10-23": "https://www.e-stat.go.jp/en/stat-search/file-download?statInfId=000040513539&fileKind=4",
    },
    "jp_gdp": {
        "2026-08-17": "https://www.esri.cao.go.jp/jp/sna/data/data_list/sokuhou/files/2026/qe262/tables/ritu-jk2621.csv",
        "2026-09-08": "https://www.esri.cao.go.jp/jp/sna/data/data_list/sokuhou/files/2026/qe262_2/tables/ritu-jk2622.csv",
    },
}
PROVIDERS = {"jp_gdp": "esri", "jp_cpi": "statjp", "jp_tokyo_cpi": "statjp", "jp_tankan": "boj", "jp_trade_balance": "mof"}
CPI_LISTINGS = {
    "2026-08-28": "https://www.e-stat.go.jp/en/stat-search/files?page=1&layout=datalist&year=20260&month=23070907&toukei=00200573&tstat=000001243876&tclass1=000001243877&cycle=1&tclass2val=0",
    "2026-10-02": "https://www.e-stat.go.jp/en/stat-search/files?page=1&layout=datalist&lid=000001490251",
    "2026-09-18": "https://www.e-stat.go.jp/en/stat-search/files?page=1&layout=datalist&lid=000001490251",
    "2026-10-23": "https://www.e-stat.go.jp/en/stat-search/files?page=1&layout=datalist&lid=000001491870",
}


def cpi_archive_listing(document, family, edition):
    """Bind the reviewed file ID to the issuer's first monthly report listing."""
    period, _, _ = reviewed(family, edition)
    page = Page(document.body).root
    tokyo = family == "jp_tokyo_cpi"
    label = "Subgroup Index for Ku-area of Tokyo" if tokyo else "Subgroup Index for Japan"
    from .bls import shift_month
    survey = shift_month(period, -1) if tokyo else period
    require("2025-Base Consumer Price Index", plain(document.body.decode("utf-8")))
    url = ARCHIVES[family][edition]
    found = []
    for group in page.find_all("ul"):
        if group.get("class") != "stat-dataset_list-detail":
            continue
        text = plain(group.text())
        for link in group.find_all("a"):
            if "https://www.e-stat.go.jp" + link.get("href") != url:
                continue
            # A later update cannot establish the original edition.
            require(re.escape(label) + r" Monthly Survey date " +
                    datetime.strptime(survey, "%Y-%m").strftime("%Y%b.") +
                    r" Update date " + edition, text)
            # Tokyo's next-month preliminary table is the second publication
            # in this monthly-report slot; this is explicitly reviewed.
            if link.get("data-release_count") != ("2" if tokyo else "1"):
                raise FetchError("e-Stat updated report publication needs review")
            found.append(url)
    if len(found) != 1:
        raise FetchError("e-Stat monthly report archive identity needs review")


def parse_cpi(document, family, since):
    from openpyxl import load_workbook
    editions = [d for d, url in ARCHIVES[family].items() if url == document.url]
    if len(editions) != 1 or not document.body.startswith(b"PK"):
        raise FetchError("e-Stat CPI edition or workbook format needs review")
    edition = editions[0]
    period, stage, _ = reviewed(family, edition)
    tokyo = family == "jp_tokyo_cpi"
    workbook = load_workbook(BytesIO(document.body), data_only=True, read_only=True)
    try:
        sheets = [s for s in workbook if s['H8'].value == 'Monthly, % Change over the year']
        if len(sheets) != 1:
            raise FetchError("e-Stat published annual-rate sheet needs review")
        sheet = sheets[0]
        label = 'Table 1-2 Subgroup Index for Ku-area of Tokyo' if tokyo else 'Table 1-1 Subgroup Index for Japan'
        if (sheet['I6'].value.strip() != label or
                sheet['I12'].value != 'All items' or sheet['J12'].value != 'All items, less fresh food'):
            raise FetchError("e-Stat CPI geography or measure headers need review")
        notes = [str(v) for row in sheet.iter_rows(values_only=True) for v in row if v is not None and 'preliminary' in str(v)]
        if tokyo and notes != ['Note) Latest monthly figures are preliminary ones.']:
            raise FetchError("e-Stat Tokyo preliminary stage needs review")
        rows = [r for r in sheet.iter_rows(values_only=True) if len(r) > 9 and
                str(r[1]).strip() == period.replace('-', '') and str(r[2]).strip() == '03']
        if len(rows) != 1 or max(str(r[1]).strip() for r in sheet.iter_rows(values_only=True)
                if len(r) > 2 and re.fullmatch(r'\d{6}', str(r[1]).strip())) != period.replace('-', ''):
            raise FetchError("e-Stat CPI headline reference period needs review")
        result = []
        pairs = [('jp_tokyo_core_cpi_yy', rows[0][9])] if tokyo else list(zip(('jp_cpi_yy', 'jp_core_cpi_yy'), rows[0][8:10]))
        for indicator, value in pairs:
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise FetchError("e-Stat CPI annual rate needs review")
            result.append(observation(document, family, edition, indicator, str(value), "xlsx_table"))
        return result if period >= since else []
    finally:
        workbook.close()


GDP_ARCHIVE = "https://www.esri.cao.go.jp/en/sna/data/sokuhou/files/2026/toukei_2026.html"
GDP_MENUS = {
    "2026-08-17": "https://www.esri.cao.go.jp/en/sna/data/sokuhou/files/2026/qe262/gdemenuea.html",
    "2026-09-08": "https://www.esri.cao.go.jp/en/sna/data/sokuhou/files/2026/qe262_2/gdemenuea.html",
}
TRADE_STYLE = "https://www.customs.go.jp/toukei/shinbun/trade-st_e/2026/d1201e.xsl"


def gdp_archive_listing(archive, menu, edition):
    """The CSV has no release stamp: bind its exact link to the dated archive."""
    _, stage, _ = reviewed("jp_gdp", edition)
    if archive.url != GDP_ARCHIVE or menu.url != GDP_MENUS.get(edition):
        raise FetchError("ESRI GDP edition listing URL needs review")
    matches = []
    for row in Page(archive.body).root.find_all("tr"):
        for link in row.find_all("a"):
            if "https://www.esri.cao.go.jp" + link.get("href") == menu.url:
                matches.append(plain(row.text()))
    ordinal = "First" if stage == "preliminary" else "Second"
    stamp = datetime.strptime(edition, "%Y-%m-%d").strftime("%b %d, %Y")
    # The issuer omits the leading zero on single-digit days.
    stamp = stamp.replace(" 0", " ")
    if len(matches) != 1:
        raise FetchError("ESRI GDP archive edition needs review")
    require(re.escape(stamp) + r" Quarterly Estimates of GDP for Apr\.-Jun\.2026 \(The " +
            ordinal + r" preliminary Estimates\)", matches[0])
    require(r"Apr\.-Jun\.2026 \(The " + ("1st" if stage == "preliminary" else "2nd") +
            r" preliminary\)", plain(menu.body.decode("utf-8")))
    links = [link for link in Page(menu.body).root.find_all("a")
             if "https://www.esri.cao.go.jp" + link.get("href") == ARCHIVES["jp_gdp"][edition]]
    if len(links) != 1 or not plain(links[0].text()).startswith(
            "Real, Seasonally Adjusted Series (Quarter-to-Quarter) (csv:"):
        raise FetchError("ESRI GDP published-rate CSV link needs review")


def parse_gdp(document, since):
    editions = [d for d, url in ARCHIVES["jp_gdp"].items() if url == document.url]
    if len(editions) != 1:
        raise FetchError("ESRI GDP CSV archive URL needs review")
    edition = editions[0]
    period, _, _ = reviewed("jp_gdp", edition)
    try:
        rows = list(csv.reader(StringIO(document.body.decode("cp932"))))
    except (UnicodeError, csv.Error) as exc:
        raise FetchError("ESRI GDP CSV format needs review") from exc
    if (len(rows) < 8 or len(rows[1]) < 2 or len(rows[5]) < 2 or rows[1][0] !=
            "Real, Seasonally Adjusted Series (Quarter-to-Quarter Percent Change)" or
            rows[5][1] != "GDP(Expenditure Approach)" or "(%)" not in rows[1]):
        raise FetchError("ESRI GDP published-rate header format needs review")
    year = None
    rates = {}
    for row in rows[7:]:
        if not row:
            continue
        stamp = re.fullmatch(r"(?:(\d{4})/ )?(1- 3|4- 6|7- 9|10-12)\.", row[0].strip())
        if not stamp:
            if row[0].strip():
                raise FetchError("ESRI GDP CSV reference period needs review")
            continue
        year = stamp[1] or year
        key = str(year) + "-Q" + str(("1- 3", "4- 6", "7- 9", "10-12").index(stamp[2]) + 1)
        if year is None or key in rates or len(row) < 2:
            raise FetchError("ESRI GDP CSV reference period needs review")
        rates[key] = row[1].strip()
    if not rates or max(rates) != period or not re.fullmatch(r"[+-]?\d+\.\d+", rates[period]):
        raise FetchError("ESRI GDP current quarter rate needs review")
    return [observation(document, "jp_gdp", edition, "jp_gdp_qq", rates[period], "csv")] if period >= since else []


def trade_stylesheet(document):
    """The publisher's XML presentation defines both units and NSA basis."""
    if document.url != TRADE_STYLE:
        raise FetchError("MOF trade stylesheet URL needs review")
    text = plain(document.body.decode("utf-8"))
    require(r"Ministry of Finance", text)
    require(r"Total \( Census, not seasonally adjusted \)", text)
    require(r"\(Unit: millions of YEN, % \)", text)


def parse_trade(document, since):
    if b"<!DOCTYPE" in document.body or b"<!ENTITY" in document.body:
        raise FetchError("MOF trade XML format needs review")
    try:
        root = ElementTree.fromstring(document.body)
    except ElementTree.ParseError as exc:
        raise FetchError("MOF trade XML format needs review") from exc
    pages = root.findall("sogakutsuki")
    if root.tag != "hodoxml" or len(pages) != 1 or pages[0].get("name") != "pg1":
        raise FetchError("MOF trade XML headline table needs review")
    page = pages[0]
    try:
        edition = datetime.strptime(page.findtext("kohyoymd", ""), "%B %d, %Y").date().isoformat()
    except ValueError as exc:
        raise FetchError("MOF trade publication date needs review") from exc
    period, stage, _ = reviewed("jp_trade_balance", edition)
    if document.url != ARCHIVES["jp_trade_balance"].get(edition):
        raise FetchError("MOF trade XML archive URL needs review")
    label = datetime.strptime(period, "%Y-%m").strftime("%B %Y")
    wording = "(Provisional)" if stage == "preliminary" else "(Exports:Detailed Data;Imports:Provisional(HS-code Basis))"
    if page.findtext("title") != "Value of Exports and Imports " + label + " " + wording or page.findtext("taishoymtonen") != label:
        raise FetchError("MOF trade reference period or stage needs review")
    tables = root.findall("chiikikunisogaku")
    if len(tables) != 1 or tables[0].get("name") != "pg2" or tables[0].findtext("taishokikan") != label:
        raise FetchError("MOF trade country table needs review")
    totals = [r for r in tables[0].findall("chiikikunisogakuinfo") if r.findtext("chiikikunikbn") == "0"]
    if len(totals) != 1 or totals[0].findtext("chiikikuni") != "Grand Total":
        raise FetchError("MOF trade national balance needs review")
    printed = totals[0].findtext("sashihikikagakue", "")
    headline = page.findtext("sashihiki/sogakutonen", "").replace("△", "-")
    if not re.fullmatch(r"[+-]?[\d,]+", printed) or printed != headline:
        raise FetchError("MOF trade printed balance disagreement needs review")
    return [observation(document, "jp_trade_balance", edition, "jp_trade_balance", printed, "structured",
                        value=float(printed.replace(',', '')) / 1000)] if period >= since else []


def parse(document, family, since="1900-01"):
    if family == "jp_trade_balance":
        return parse_trade(document, since)
    if family == "jp_tankan":
        from pypdf import PdfReader
        if not document.body.startswith(b"%PDF"):
            raise FetchError("BOJ Tankan PDF format needs review")
        body = re.sub(r"\s+", " ", PdfReader(BytesIO(document.body)).pages[0].extract_text())
        stamp = require(r"Tankan Summary \((June|September|December) 2026\) \d+th Short-Term Economic Survey of Enterprises in Japan (\w+ \d{1,2}, 2026)", body)
        edition = datetime.strptime(stamp[2], "%B %d, %Y").date().isoformat()
        period, _, _ = reviewed(family, edition)
        if document.url != ARCHIVES[family].get(edition) or stamp[1] != {"Q2": "June", "Q3": "September", "Q4": "December"}[period[-2:]]:
            raise FetchError("BOJ Tankan survey period or URL needs review")
        require(r'1\. Business Conditions DI \("favorable" - "unfavorable"\), % points Large Enterprises Medium-sized Enterprises Small Enterprises', body)
        # Large enterprises occupy the first six columns: previous actual,
        # previous forecast, current actual, change, forecast, change.
        row = require(r"Changes Changes Changes Changes Changes Changes Manufacturing ([+-]?\d+) ([+-]?\d+) ([+-]?\d+) ([+-]?\d+) ([+-]?\d+) ([+-]?\d+)", body)
        return [observation(document, family, edition, "jp_tankan_large_manufacturers_di", row[3], "pdf_table")] if period >= since else []
    if family in {"jp_cpi", "jp_tokyo_cpi"}:
        return parse_cpi(document, family, since)
    if family == "jp_gdp":
        return parse_gdp(document, since)
    raise FetchError("Japan issuer family needs review")


def fetch(client, family, since):
    if family == "jp_gdp":
        archive = None
        for edition, (period, _, _) in due(EDITIONS[family], family=family).items():
            if period >= since:
                if edition not in GDP_MENUS:
                    raise FetchError("ESRI GDP due archive listing needs review")
                archive = archive or client.get("esri", GDP_ARCHIVE)
                gdp_archive_listing(archive, client.get("esri", GDP_MENUS[edition]), edition)
    if family == "jp_trade_balance" and any(period >= since for period, _, _ in due(EDITIONS[family], family=family).values()):
        trade_stylesheet(client.get("mof", TRADE_STYLE))
    if family in {"jp_cpi", "jp_tokyo_cpi"}:
        for edition, (period, _, _) in due(EDITIONS[family], family=family).items():
            if period >= since:
                if edition not in CPI_LISTINGS:
                    raise FetchError("e-Stat due CPI report archive needs review")
                cpi_archive_listing(client.get("statjp", CPI_LISTINGS[edition]), family, edition)
    return fetch_archives(client, family, since, PROVIDERS[family], ARCHIVES[family], parse)
